import streamlit as st
import openpyxl
import pandas as pd
import io
import zipfile
import base64
from datetime import date, datetime
from template_data import TEMPLATE_B64

def load_template():
    return io.BytesIO(base64.b64decode(TEMPLATE_B64))

# 商品マスタ（テンプレートV6:W11より）
PRICE_MASTER = {
    "ツナプレーン": 520,
    "ツナおかず": 580,
    "ツナラー油": 580,
    "ツナ味噌": 580,
    "ツナガリ": 580,
    "ツナカレー": 580,
}
TAX_RATE = 0.08
ITEM_START_ROW = 18
ITEM_END_ROW = 33

st.set_page_config(page_title="Invoice Generator | 株式会社 新家", layout="centered")

st.markdown("""
<style>
    /* ベース */
    html, body, [class*="css"], .stApp {
        font-family: 'Hiragino Kaku Gothic Pro', 'Noto Sans JP', sans-serif;
        background-color: #0a0a0a !important;
        color: rgba(255,255,255,0.9);
        line-height: 1.65;
    }

    /* Streamlit内部の白背景を全て上書き */
    .main, .block-container {
        background-color: #0a0a0a !important;
    }
    section[data-testid="stSidebar"] {
        background-color: #0f0f0f !important;
    }

    /* ヘッダー */
    .app-header {
        padding: 3rem 0 1.8rem 0;
        border-bottom: 1px solid rgba(255,255,255,0.08);
        margin-bottom: 2.5rem;
    }
    .app-header h1 {
        font-size: 1.3rem;
        font-weight: 700;
        letter-spacing: 0.08em;
        color: rgba(255,255,255,0.95);
        margin: 0;
    }
    .app-header p {
        font-size: 0.78rem;
        color: rgba(255,255,255,0.4);
        margin: 0.4rem 0 0 0;
        letter-spacing: 0.04em;
    }

    /* セクションラベル */
    .section-label {
        font-size: 0.65rem;
        font-weight: 700;
        letter-spacing: 0.15em;
        text-transform: uppercase;
        color: rgba(255,255,255,0.35);
        margin-bottom: 0.5rem;
    }

    /* ボタン */
    .stButton > button {
        background-color: #ffffff !important;
        color: #0a0a0a !important;
        border: none !important;
        border-radius: 6px;
        font-size: 0.85rem;
        font-weight: 600;
        letter-spacing: 0.06em;
        padding: 0.7rem 1.5rem;
        transition: opacity 0.15s ease;
    }
    .stButton > button:hover {
        background-color: rgba(255,255,255,0.88) !important;
        color: #0a0a0a !important;
    }

    /* プライマリボタン */
    .stButton > button[kind="primary"] {
        background-color: #ffffff !important;
        color: #0a0a0a !important;
    }

    /* 入力フィールド */
    input, textarea, [data-testid="stTextInput"] input {
        background-color: #141414 !important;
        color: rgba(255,255,255,0.9) !important;
        border: 1px solid rgba(255,255,255,0.1) !important;
        border-radius: 6px !important;
    }
    input:focus, textarea:focus {
        border-color: rgba(255,255,255,0.3) !important;
        box-shadow: none !important;
    }

    /* ラベルテキスト */
    label, .stTextInput label {
        color: rgba(255,255,255,0.5) !important;
        font-size: 0.78rem !important;
    }

    /* ファイルアップローダー */
    [data-testid="stFileUploader"] {
        background-color: #141414 !important;
        border: 1px dashed rgba(255,255,255,0.15) !important;
        border-radius: 8px !important;
    }
    [data-testid="stFileUploader"] *,
    [data-testid="stFileUploader"] p,
    [data-testid="stFileUploader"] span,
    [data-testid="stFileUploader"] small,
    [data-testid="stFileUploader"] div,
    [data-testid="stFileUploaderDropzone"] *,
    [data-testid="stFileUploaderDropzone"] {
        background-color: #141414 !important;
        color: rgba(255,255,255,0.65) !important;
    }
    [data-testid="stFileUploaderDropzone"] {
        border: none !important;
    }
    /* アップロード済みファイル表示 */
    [data-testid="stFileUploaderFile"] {
        background-color: #1e1e1e !important;
        border: 1px solid rgba(255,255,255,0.08) !important;
        border-radius: 6px !important;
    }
    [data-testid="stFileUploaderFile"] * {
        color: rgba(255,255,255,0.75) !important;
    }
    /* Browse files ボタン */
    [data-testid="stFileUploader"] button {
        background-color: #2a2a2a !important;
        color: rgba(255,255,255,0.8) !important;
        border: 1px solid rgba(255,255,255,0.15) !important;
        border-radius: 6px !important;
    }

    /* expander */
    [data-testid="stExpander"] {
        background-color: #141414 !important;
        border: 1px solid rgba(255,255,255,0.07) !important;
        border-radius: 8px !important;
    }
    [data-testid="stExpander"] summary {
        color: rgba(255,255,255,0.85) !important;
    }

    /* dataframe / テーブル */
    [data-testid="stDataFrame"], .stDataFrame {
        background-color: #141414 !important;
        border: 1px solid rgba(255,255,255,0.07) !important;
        border-radius: 6px !important;
    }

    /* メトリクス */
    [data-testid="stMetric"] {
        background-color: #141414;
        border: 1px solid rgba(255,255,255,0.07);
        border-radius: 8px;
        padding: 1rem;
    }
    [data-testid="stMetricLabel"] {
        color: rgba(255,255,255,0.45) !important;
        font-size: 0.72rem !important;
        letter-spacing: 0.06em;
    }
    [data-testid="stMetricValue"] {
        color: rgba(255,255,255,0.95) !important;
        font-size: 1.1rem !important;
        font-weight: 600;
    }

    /* caption / info テキスト */
    .stCaption, [data-testid="stCaptionContainer"] {
        color: rgba(255,255,255,0.35) !important;
    }

    /* warning / error / success */
    [data-testid="stAlert"] {
        background-color: #1a1a1a !important;
        border: 1px solid rgba(255,255,255,0.1) !important;
        border-radius: 6px !important;
        color: rgba(255,255,255,0.8) !important;
    }

    /* progress bar */
    [data-testid="stProgressBar"] > div {
        background-color: rgba(255,255,255,0.15) !important;
        border-radius: 4px;
    }
    [data-testid="stProgressBar"] > div > div {
        background-color: #ffffff !important;
    }

    /* divider */
    hr {
        border-color: rgba(255,255,255,0.06) !important;
    }

    /* フッター非表示 */
    footer { visibility: hidden; }
    #MainMenu { visibility: hidden; }
</style>

<div class="app-header">
    <h1>Invoice Generator</h1>
    <p>株式会社 新家 / 納品書から請求書を自動生成します</p>
</div>
""", unsafe_allow_html=True)


def read_nouhinshо(file_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    ws = wb.active

    store_name = ws["A3"].value or ""
    delivery_date = ws["N4"].value

    price_map = {}
    for row in ws.iter_rows(min_row=6, max_row=11, min_col=22, max_col=23, values_only=True):
        if row[0] and row[1]:
            price_map[str(row[0])] = int(row[1])
    for k, v in PRICE_MASTER.items():
        price_map.setdefault(k, v)

    items = []
    for r in range(18, 30):
        name = ws.cell(row=r, column=2).value
        qty  = ws.cell(row=r, column=10).value
        if name and qty and int(qty) > 0:
            unit_price = price_map.get(str(name), 0)
            items.append({
                "name": str(name),
                "unit_price": unit_price,
                "quantity": int(qty),
                "amount": int(qty) * unit_price,
            })

    if not store_name.strip() or not items:
        return None

    subtotal = sum(i["amount"] for i in items)
    return {
        "store_name": store_name.strip(),
        "delivery_date": delivery_date,
        "items": items,
        "subtotal": subtotal,
        "tax": int(subtotal * TAX_RATE),
        "total": subtotal + int(subtotal * TAX_RATE),
    }


def create_invoice(store_name, dated_groups, billing_month, billing_subject):
    """お父さんのテンプレートをコピーして数字だけ書き換える
    dated_groups: [{"date": datetime, "items": [...]}, ...]  日付ごとのグループ
    """
    wb = openpyxl.load_workbook(load_template())
    ws = wb.active

    # 開いた時に必ず数式を再計算させる（合計額が反映されないバグの修正）
    wb.calculation.fullCalcOnLoad = True

    # 宛先・日付・件名
    ws["A3"] = store_name
    ws["N4"] = date.today()

    # 件名（例: 令和8年5月分味噌加工品代金）
    reiwa = date.today().year - 2018
    ws["C6"] = f"令和{reiwa}年{billing_month}{billing_subject}"

    # 明細を一度クリア（行18〜33のA・B・J・K列）
    for r in range(ITEM_START_ROW, ITEM_END_ROW + 1):
        ws.cell(row=r, column=1).value = None   # A: 日付
        ws.cell(row=r, column=2).value = None   # B: 商品名
        ws.cell(row=r, column=10).value = None  # J: 数量
        ws.cell(row=r, column=11).value = None  # K: 単位

    # 日付グループごとに明細を書き込む
    current_row = ITEM_START_ROW
    for group in dated_groups:
        if current_row > ITEM_END_ROW:
            break

        # 日付ラベル（例: "5月4日"）
        d = group["date"]
        date_label = f"{d.month}月{d.day}日" if isinstance(d, datetime) else ""

        # 同じ日付内で同じ商品は合算
        merged = {}
        for item in group["items"]:
            k = item["name"]
            if k in merged:
                merged[k]["quantity"] += item["quantity"]
                merged[k]["amount"]   += item["amount"]
            else:
                merged[k] = dict(item)

        # この日付のアイテムを書き込む
        first_in_group = True
        for item in merged.values():
            if current_row > ITEM_END_ROW:
                break
            if first_in_group:
                ws.cell(row=current_row, column=1).value = date_label  # A列: 納品日（各グループ先頭のみ）
                first_in_group = False
            ws.cell(row=current_row, column=2).value  = item["name"]     # B列: 商品名
            ws.cell(row=current_row, column=10).value = item["quantity"]  # J列: 数量
            ws.cell(row=current_row, column=11).value = "個"             # K列: 単位
            current_row += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- UI ----

reiwa_now = date.today().year - 2018

col1, col2 = st.columns(2)
with col1:
    st.markdown('<p class="section-label">請求月</p>', unsafe_allow_html=True)
    billing_month = st.text_input(
        "請求月",
        value=f"{date.today().month}月分",
        label_visibility="collapsed"
    )
with col2:
    st.markdown('<p class="section-label">件名（商品・サービス名）</p>', unsafe_allow_html=True)
    billing_subject = st.text_input(
        "件名",
        value="味噌加工品代金",
        label_visibility="collapsed"
    )

st.caption(f"請求書の件名: 令和{reiwa_now}年{billing_month}{billing_subject}")

st.write("")
st.markdown('<p class="section-label">納品書ファイル</p>', unsafe_allow_html=True)
uploaded_files = st.file_uploader(
    "納品書ファイル",
    type=["xlsx", "xls"],
    accept_multiple_files=True,
    label_visibility="collapsed"
)

if uploaded_files:
    st.caption(f"{len(uploaded_files)} 件のファイルを読み込みました")

st.write("")

if st.button("請求書を生成する", type="primary", use_container_width=True):
    if not uploaded_files:
        st.warning("納品書ファイルをアップロードしてください")
        st.stop()

    store_data = {}
    errors = []

    progress = st.progress(0, text="読み取り中...")
    for i, f in enumerate(uploaded_files):
        progress.progress((i + 1) / len(uploaded_files), text=f"読み取り: {f.name}")
        result = read_nouhinshо(f.read())
        if result is None:
            errors.append(f.name)
            continue
        store = result["store_name"]
        if store not in store_data:
            store_data[store] = {"dated_groups": []}
        # 日付とアイテムをセットで保持（日別詳細を維持するため）
        store_data[store]["dated_groups"].append({
            "date":  result["delivery_date"],
            "items": result["items"],
        })
    progress.empty()

    if errors:
        st.warning(f"読み取れなかったファイル: {', '.join(errors)}")
    if not store_data:
        st.error("データを抽出できませんでした")
        st.stop()

    st.write("")
    st.markdown('<p class="section-label">内容確認</p>', unsafe_allow_html=True)
    st.caption("金額に誤りがないか確認してからダウンロードしてください")

    all_ok = True
    for store_name, data in store_data.items():
        # 全アイテムを合算（金額確認・合計表示用）
        all_items_flat = [item for g in data["dated_groups"] for item in g["items"]]
        merged_all = {}
        for item in all_items_flat:
            k = item["name"]
            if k in merged_all:
                merged_all[k]["quantity"] += item["quantity"]
                merged_all[k]["amount"]   += item["amount"]
            else:
                merged_all[k] = dict(item)
        merged_items = list(merged_all.values())

        subtotal = sum(i["amount"] for i in merged_items)
        tax      = int(subtotal * TAX_RATE)
        total    = subtotal + tax

        # 単価×数量の検証
        check_errors = [
            f"{i['name']}: 計算値¥{i['unit_price']*i['quantity']:,} ≠ ¥{i['amount']:,}"
            for i in merged_items if i["unit_price"] * i["quantity"] != i["amount"]
        ]

        status = "-- 確認済み" if not check_errors else "-- 要確認"
        with st.expander(f"{store_name}　{status}　合計 ¥{total:,}（税込）", expanded=True):
            if check_errors:
                st.error("金額の不一致が検出されました: " + " / ".join(check_errors))
                all_ok = False

            # 日別明細を表示
            for group in data["dated_groups"]:
                d = group["date"]
                label = f"{d.month}月{d.day}日" if isinstance(d, datetime) else "日付不明"
                st.caption(f"**{label}**")
                df = pd.DataFrame(group["items"])[["name", "unit_price", "quantity", "amount"]]
                df.columns = ["商品名", "単価", "数量", "金額"]
                df["単価"] = df["単価"].map(lambda x: f"¥{x:,}")
                df["金額"] = df["金額"].map(lambda x: f"¥{x:,}")
                st.dataframe(df, use_container_width=True, hide_index=True)

            c1, c2, c3 = st.columns(3)
            c1.metric("小計（税抜）", f"¥{subtotal:,}")
            c2.metric("消費税（8%）", f"¥{tax:,}")
            c3.metric("合計（税込）", f"¥{total:,}")

    st.write("")

    if not all_ok:
        st.error("金額の不一致があります。納品書ファイルを確認してください。")
    else:
        zip_buf = io.BytesIO()
        with zipfile.ZipFile(zip_buf, "w") as zf:
            for store_name, data in store_data.items():
                excel_bytes = create_invoice(
                    store_name,
                    data["dated_groups"],
                    billing_month,
                    billing_subject,
                )
                safe = store_name.replace("/", "_").replace(" ", "_").replace("　", "_")
                zf.writestr(f"令和{date.today().year-2018}年{billing_month}請求書（{safe}）.xlsx", excel_bytes)

        zip_buf.seek(0)
        st.download_button(
            label=f"請求書をダウンロード（{len(store_data)}店舗分 / ZIP）",
            data=zip_buf,
            file_name=f"請求書一括_{billing_month}.zip",
            mime="application/zip",
            use_container_width=True,
            type="primary",
        )
